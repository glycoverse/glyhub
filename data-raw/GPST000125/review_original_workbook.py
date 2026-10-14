"""Verify the original workbook and record targeted native-vector evidence."""
from pathlib import Path
import zipfile,xml.etree.ElementTree as E,json,hashlib,collections,math
import openpyxl
D=Path(__file__).resolve().parent
p=D/'mmc5.xlsx';b=p.read_bytes();h=hashlib.sha256(b).hexdigest()
prior=json.loads((D/'source-cell-records.json').read_text())
assert len(b)==prior['original_byte_size']==253603
assert h==prior['original_sha256']=='ae008af32e54980a6add4907da91a49723f9affb710c32e00f672c1e8e95ef41'
wb=openpyxl.load_workbook(p,data_only=True,read_only=True)
ss=json.loads((D/'matrix-specs.json').read_text());matches=0
for s in ss:
 for i,r in enumerate(s['row_indices']):
  for j,c in enumerate(s['column_indices']):
   assert abs(wb[s['sheet']].cell(r,c).value-s['raw_matrix'][i][j])<1e-12
   matches+=1
z=zipfile.ZipFile(p);ns={'xdr':'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
root=E.fromstring(z.read('xl/drawings/drawing1.xml'))
labels={('ellipse','33CC33'):'Man',('ellipse','FFFF00'):'Gal',('ellipse','0000FF'):'Glc',('rect','0000FF'):'GlcNAc',('rect','FFFF00'):'GalNAc',('triangle','FF0000'):'Fuc',('rect','A54399'):'Neu5Ac'}
records=[]
for a in root:
 row=int(a.find('xdr:from/xdr:row',ns).text)+1
 if row not in [18,19,20,28,41,75]:continue
 symbols=[];counts=collections.Counter()
 for s in a.findall('.//xdr:sp',ns):
  sp=s.find('xdr:spPr',ns);g=sp.find('a:prstGeom',ns);fill=sp.find('a:solidFill/a:srgbClr',ns)
  if g is None or fill is None:continue
  key=(g.get('prst'),fill.get('val'))
  if key not in labels:continue
  rot=sp.find('a:xfrm',ns).get('rot','0')
  sym={'shape':key[0],'fill_srgb':key[1],'rotation_60000ths_degree':int(rot),'source_symbol':labels[key]}
  symbols.append(sym);counts[labels[key]]+=1
 generic={'Hex':sum(counts[x] for x in ['Man','Gal','Glc']),'HexNAc':sum(counts[x] for x in ['GlcNAc','GalNAc']),'dHex':counts['Fuc'],'Neu5Ac':counts['Neu5Ac']}
 comp=''.join(k+str(generic[n]) for k,n in [('H','Hex'),('N','HexNAc'),('F','dHex'),('S','Neu5Ac')] if generic[n])
 records.append({'sheet':'Table_S4_1','source_row':row,'source_feature_number':wb['Table_S4_1'].cell(row,1).value,'source_name':wb['Table_S4_1'].cell(row,2).value,'source_diagram_cell':f'C{row}','drawing_part':'xl/drawings/drawing1.xml','source_diagram_sha256':hashlib.sha256(E.tostring(a)).hexdigest(),'native_vector_symbols':symbols,'specific_symbol_counts':dict(counts),'generic_counts':generic,'diagram_composition':comp,'manual_render_reviewed':True,'rendered_pdf_page_1based':5 if row==75 else (3 if row==41 else 2),'decision':'unresolved_source_name_component_structure_and_diagram_disagreement' if row==20 else 'composition_supported_by_source_name_and_diagram_despite_inconsistent_registry_id','canonical_composition':None if row==20 else comp,'linkages_inferred_from_diagram':False})
records.sort(key=lambda x:x['source_row'])
result={'source_filename':'mmc5.xlsx','source_size_bytes':len(b),'source_sha256':h,'source_md5':hashlib.md5(b).hexdigest(),'matches_prior_original_file_hash':True,'quantitative_cells_checked':matches,'original_bytes_retained':True,'drawings_preserved':True,'review_scope':'Six previously conflicting rows only; native vector shapes were rendered with LibreOffice and visually inspected. Symbol counts were independently checked from DrawingML. No new anomeric or linkage assignment is inferred.','snfg_mapping_note':'Colors/shapes follow the article SNFG convention. Rotated purple rectangles render as Neu5Ac diamonds; red triangles are Fuc. Source names corroborate the residue labels.','records':records}
(D/'original-workbook-review.json').write_text(json.dumps(result,indent=2,ensure_ascii=False)+'\n')
print(json.dumps({'bytes':len(b),'sha256':h,'quantitative_cells_checked':matches,'records':[{k:v for k,v in x.items() if k in ['source_row','source_feature_number','diagram_composition','canonical_composition','specific_symbol_counts']} for x in records]},indent=2))
