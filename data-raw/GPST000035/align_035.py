import json,re,statistics,csv
from pathlib import Path
import openpyxl
P=Path(__file__).resolve().parent
x=json.load(open(P/'matrix-specs.json'))
a=openpyxl.load_workbook(P/'GPST000035_media_MOESM4.xlsx',data_only=True)
b=openpyxl.load_workbook(P/'GPST000035_media_MOESM5.xlsx',data_only=True)['S1']
def norm(t):return re.sub('[^a-z0-9]','',t.lower())
cols={norm(b.cell(3,c).value):c for c in range(36,149,4)}
rows=[r for r in range(5,183) if isinstance(b.cell(r,1).value,(int,float))]
out=[]
for s in x:
 sh=a[s['sheet']]; panel=re.sub(r'^S\d+\s*','',s['sheet']);c=cols.get(norm(panel))
 for i,r in enumerate(s['row_indices']):
  rt=[sh.cell(r,j-1).value for j in s['column_indices']];rel=[sh.cell(r,j+2).value for j in s['column_indices']]
  rtmean=statistics.mean(t for t in rt if isinstance(t,(int,float))); relmean=statistics.mean(t for t in rel if isinstance(t,(int,float)))
  matches=[]; near=[]
  if c:
   for rr in rows:
    br,ba=b.cell(rr,c).value,b.cell(rr,c+1).value
    if isinstance(br,(int,float)) and isinstance(ba,(int,float)):
     d=abs(br-rtmean); e=abs(ba-relmean)
     if d<=.011 and e<=.011:matches.append(rr)
     if d<=.2:near.append((rr,b.cell(rr,1).value,b.cell(rr,12).value,br,ba,e))
  out.append(dict(sheet=s['sheet'],source_row=r,feature_id=s['feature_ids'][i],source_composition=s['row_annotations'][i]['source_composition'],rt_mean=rtmean,area_mean=relmean,summary_column=c,matches=matches,glycan_number=b.cell(matches[0],1).value if len(matches)==1 else None,global_composition=b.cell(matches[0],12).value if len(matches)==1 else None,near=near if len(matches)!=1 else []))
json.dump(out,open(P/'feature_alignment_audit.json','w'),indent=2)
from collections import Counter
print('mapping',Counter('QC' if z['summary_column'] is None else len(z['matches']) for z in out))
for z in out:
 if z['summary_column'] and len(z['matches'])!=1:print(z)
# Only author-number mappings supported by both composition and within-panel
# averaged RT/relative abundance are automatically approved. No nearest-RT guessing.
for z in out:
 z['alignment_status']='unresolved'
 z['aligned_feature_id']=z['feature_id']
 if len(z['matches'])==1 and re.sub(r'\([a-z]\)','',z['global_composition'])==z['source_composition']:
  z['alignment_status']='verified_author_key'
  z['aligned_feature_id']='GPST000035_glycan_'+str(z['glycan_number'])
 elif not z['summary_column']:z['alignment_status']='QC_not_in_global_table'
 z['rt_composition_candidates']=[v for v in z['near'] if abs(v[3]-z['rt_mean'])<=.011 and re.sub(r'\([a-z]\)','',v[2])==z['source_composition']]
json.dump(out,open(P/'feature_alignment_audit.json','w'),indent=2)
fields=['sheet','source_row','feature_id','aligned_feature_id','alignment_status','glycan_number','global_composition','source_composition','rt_mean','area_mean','summary_column']
with open(P/'feature_alignment.csv','w') as f:
 w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');w.writeheader();w.writerows(out)
print('APPROVED',sum(z['alignment_status']=='verified_author_key' for z in out),'FINAL_ROWS',len(set(z['aligned_feature_id'] for z in out)))
print('unresolved with exact RT+composition only',sum(len(z['rt_composition_candidates'])==1 for z in out))
